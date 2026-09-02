#!/usr/bin/env python3
"""Generate the deterministic parity input corpus under ``fixtures/inputs/``.

Constitution III wants goldens captured from the Python reference rather than hand-written
expectations. That only works if the *inputs* are fixed too, so this script builds every
synthetic case from a seed: the same run on any machine produces byte-identical files.

Cases that came from real editorial documents (``word-brsm``, ``word-den-znaniy``) are left
alone -- they are checked in as-is and this script never touches them.

Run through ``tools/capture-reference`` (``cargo run -p capture-reference -- --inputs``), or
directly with the harness virtualenv:

    tools/capture-reference/.cache/venv/bin/python tools/capture-reference/make_inputs.py
"""

from __future__ import annotations

import copy
import shutil
import sys
import zipfile
from pathlib import Path

from PIL import Image, ImageDraw

REPO_ROOT = Path(__file__).resolve().parents[2]
INPUTS = REPO_ROOT / "fixtures" / "inputs"

# Cases captured from real editorial documents. Never regenerated.
PRESERVED = {"word-brsm", "word-den-znaniy"}

# A fixed palette, so an image's bytes depend only on its index.
PALETTE = [
    (198, 92, 61),
    (61, 108, 198),
    (86, 158, 96),
    (191, 160, 58),
    (132, 78, 172),
    (58, 160, 172),
    (172, 58, 118),
    (100, 116, 139),
]


def solid_image(width: int, height: int, index: int) -> Image.Image:
    """A deterministic image whose content is a pure function of its index and size."""
    base = PALETTE[index % len(PALETTE)]
    image = Image.new("RGB", (width, height), base)
    draw = ImageDraw.Draw(image)
    # A few shapes so the encoder has something to work with and JPEG quality steps
    # actually change the output size.
    for step in range(0, min(width, height) // 2, 37):
        shade = ((base[0] + step) % 256, (base[1] + step * 3) % 256, (base[2] + step * 7) % 256)
        draw.rectangle([step, step, width - step, height - step], outline=shade, width=3)
    draw.rectangle([width // 4, height // 8, 3 * width // 4, height // 2], fill=PALETTE[(index + 3) % len(PALETTE)])
    return image


def noisy_image(width: int, height: int, index: int) -> Image.Image:
    """A deterministic but incompressible image, for the size-budget case.

    A flat colour compresses far below the byte budget no matter the quality, which would
    make the quality-step search a no-op. This pattern does not.
    """
    image = Image.new("RGB", (width, height))
    pixels = image.load()
    seed = 1103515245 * (index + 1) + 12345
    state = seed & 0xFFFFFFFF
    for y in range(height):
        for x in range(width):
            state = (1103515245 * state + 12345) & 0xFFFFFFFF
            pixels[x, y] = ((state >> 16) & 0xFF, (state >> 8) & 0xFF, state & 0xFF)
    return image


def write_image(path: Path, image: Image.Image) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    suffix = path.suffix.lower()
    if suffix in {".jpg", ".jpeg"}:
        image.save(path, format="JPEG", quality=92, optimize=True, progressive=True)
    elif suffix == ".png":
        image.save(path, format="PNG", optimize=True, compress_level=9)
    elif suffix == ".webp":
        image.save(path, format="WEBP", quality=92, method=6)
    elif suffix == ".gif":
        image.convert("P", palette=Image.Palette.ADAPTIVE).save(path, format="GIF")
    else:
        raise ValueError(f"unhandled image suffix: {suffix}")


def reset_case(name: str) -> Path:
    case = INPUTS / name
    if case.exists():
        shutil.rmtree(case)
    case.mkdir(parents=True)
    return case


def write_text(path: Path, body: str) -> None:
    """Writes UTF-8 with LF endings, so goldens do not depend on the checkout's platform."""
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(body, encoding="utf-8", newline="\n")


# --------------------------------------------------------------------------------------
# Plain-text and Markdown cases
# --------------------------------------------------------------------------------------


def case_markers() -> None:
    """All four markers in one document, plus the clearing behaviour after a float."""
    case = reset_case("markers")
    for idx in range(1, 6):
        write_image(case / "images" / f"photo{idx}.jpg", solid_image(900, 600, idx))
    write_text(
        case / "input.txt",
        "Marker Coverage\n"
        "\n"
        "The lead paragraph introduces the item.\n"
        "\n"
        "[image:1]\n"
        "\n"
        "A body paragraph follows the full-width photo.\n"
        "\n"
        "[images:2,3]\n"
        "\n"
        "Another body paragraph after the row.\n"
        "\n"
        "[image-left:4]\n"
        "\n"
        "Text that wraps around the left float and keeps going for long enough to matter.\n"
        "\n"
        "[image-right:5]\n"
        "\n"
        "Text that wraps around the right float.\n"
        "\n"
        "[image:1]\n"
        "\n"
        "A closing paragraph, which forces the float to clear before the full-width photo.\n",
    )


def case_single_markers() -> None:
    """Each marker alone, so a golden isolates one layout's element structure."""
    for name, marker in (
        ("marker-full-width", "[image:1]"),
        ("marker-row", "[images:1,2]"),
        ("marker-float-left", "[image-left:1]"),
        ("marker-float-right", "[image-right:1]"),
    ):
        case = reset_case(name)
        for idx in range(1, 3):
            write_image(case / "images" / f"photo{idx}.jpg", solid_image(900, 600, idx))
        write_text(
            case / "input.txt",
            f"Single Layout\n\nOne paragraph before.\n\n{marker}\n\nOne paragraph after.\n",
        )


def case_row_of_one() -> None:
    """``[images:1]`` -- a row marker holding a single photo.

    The reference does *not* normalise this to a full-width placement; only ``image``,
    ``image-left`` and ``image-right`` require exactly one index. data-model.md's INV-4 claims
    the opposite, so this case is what settles it.
    """
    case = reset_case("row-of-one")
    write_image(case / "images" / "photo1.jpg", solid_image(900, 600, 1))
    write_text(case / "input.txt", "Row Of One\n\nBefore.\n\n[images:1]\n\nAfter.\n")


def case_cyrillic_title() -> None:
    """The transliteration table, including the Belarusian and Ukrainian rows."""
    case = reset_case("cyrillic-title")
    write_image(case / "images" / "photo1.jpg", solid_image(900, 600, 2))
    write_text(
        case / "input.txt",
        "День Конституции Республики Беларусь\n"
        "\n"
        "Абзац с щедрым объёмом текста, ъ и ь, ёлки и юмор.\n"
        "\n"
        "[image:1]\n"
        "\n"
        "Заключительный абзац.\n",
    )


def case_cyrillic_variants() -> None:
    """Titles chosen to exercise every distinct row of ``CYRILLIC_TRANSLIT``."""
    titles = {
        "translit-lower": "абвгдеёжзийклмнопрстуфхцчшщъыьэюя",
        "translit-upper": "АБВГДЕЁЖЗИЙКЛМНОПРСТУФХЦЧШЩЪЫЬЭЮЯ",
        "translit-belarusian-ukrainian": "і ї є ў Іван Їжак Європа Ўзвышша",
        "translit-punctuation": "Hello & Goodbye -- 2026 (draft) #1!",
        "translit-empty": "!!! ??? ...",
    }
    for name, title in titles.items():
        case = reset_case(name)
        write_image(case / "images" / "photo1.jpg", solid_image(600, 400, 4))
        write_text(case / "input.txt", f"{title}\n\nBody paragraph.\n\n[image:1]\n")


def case_natural_sort() -> None:
    """``photo1`` .. ``photo12``: bytewise ordering would put 10 before 2."""
    case = reset_case("natural-sort")
    for idx in range(1, 13):
        write_image(case / "images" / f"photo{idx}.jpg", solid_image(400, 300, idx))
    markers = "\n\n".join(f"[image:{idx}]" for idx in range(1, 13))
    write_text(case / "input.txt", f"Natural Sort\n\nLead paragraph.\n\n{markers}\n")


def case_missing_photo() -> None:
    """A marker naming an index past the end of the image list.

    The reference raises ``ValueError`` here; the port downgrades it to a warning (FR-034), so
    the golden captured for this case is the reference's *error text*, not a fragment.
    """
    case = reset_case("missing-photo")
    write_image(case / "images" / "photo1.jpg", solid_image(600, 400, 1))
    write_text(case / "input.txt", "Missing Photo\n\nLead.\n\n[image:1]\n\n[image:7]\n\nTail.\n")


def case_unused_photo() -> None:
    """More photos than markers: the surplus is warned about and not uploaded."""
    case = reset_case("unused-photo")
    for idx in range(1, 4):
        write_image(case / "images" / f"photo{idx}.jpg", solid_image(600, 400, idx))
    write_text(case / "input.txt", "Unused Photo\n\nLead.\n\n[image:1]\n\nTail.\n")


def case_oversized() -> None:
    """A photo that forces the quality-step search down from 85."""
    case = reset_case("oversized")
    write_image(case / "images" / "photo1.jpg", noisy_image(2400, 1800, 1))
    write_text(case / "input.txt", "Oversized\n\nLead.\n\n[image:1]\n")


def case_formats() -> None:
    """One case per input format, which is what pins the published extension rule."""
    case = reset_case("formats")
    write_image(case / "images" / "photo1.jpg", solid_image(800, 600, 1))
    write_image(case / "images" / "photo2.png", solid_image(800, 600, 2))
    write_image(case / "images" / "photo3.webp", solid_image(800, 600, 3))
    write_image(case / "images" / "photo4.gif", solid_image(800, 600, 4))
    markers = "\n\n".join(f"[image:{idx}]" for idx in range(1, 5))
    write_text(case / "input.txt", f"Formats\n\nLead.\n\n{markers}\n")


def case_mixed_orientation() -> None:
    """Portrait and landscape interleaved, for arrangement's shape rule."""
    case = reset_case("mixed-orientation")
    sizes = [(900, 600), (600, 900), (900, 600), (600, 900), (800, 800), (1200, 500)]
    for idx, (width, height) in enumerate(sizes, start=1):
        write_image(case / "images" / f"photo{idx}.jpg", solid_image(width, height, idx))
    write_text(
        case / "input.txt",
        "Mixed Orientation\n\nLead.\n\n[images:1,2]\n\nMiddle.\n\n[image-right:3]\n\nTail.\n",
    )


def case_portraits() -> None:
    """Twenty portraits for the SC-003 headroom benchmark.

    A pale head disc sits in the upper third of every frame, at a position recorded in
    ``heads.json`` so the Rust benchmark can assert the head survives the default crop.
    """
    case = reset_case("portraits")
    import json

    heads = []
    for idx in range(1, 21):
        width = 600 + (idx % 4) * 100
        height = 900 + (idx % 5) * 80
        image = solid_image(width, height, idx)
        draw = ImageDraw.Draw(image)
        head_cx = width // 2
        head_cy = int(height * (0.14 + 0.04 * (idx % 4)))
        head_r = int(min(width, height) * 0.11)
        draw.ellipse(
            [head_cx - head_r, head_cy - head_r, head_cx + head_r, head_cy + head_r],
            fill=(245, 224, 205),
            outline=(60, 40, 25),
            width=4,
        )
        name = f"portrait{idx:02d}.jpg"
        write_image(case / "images" / name, image)
        heads.append(
            {
                "file_name": name,
                "width": width,
                "height": height,
                "head": {
                    "x": head_cx - head_r,
                    "y": head_cy - head_r,
                    "width": head_r * 2,
                    "height": head_r * 2,
                },
            }
        )
    write_text(case / "heads.json", json.dumps(heads, indent=2, ensure_ascii=False) + "\n")
    write_text(case / "input.txt", "Portraits\n\nLead.\n\n[image:1]\n")


def case_markdown_title() -> None:
    """A Markdown ``#`` heading, which the reference treats as the title."""
    case = reset_case("markdown-title")
    write_image(case / "images" / "photo1.jpg", solid_image(700, 500, 5))
    write_text(
        case / "input.md",
        "# Markdown Heading\n\nLead paragraph.\n\n[image:1]\n\nTail paragraph.\n",
    )


def case_escaping() -> None:
    """Characters the renderer must escape, in the title and in the body."""
    case = reset_case("escaping")
    write_image(case / "images" / "photo1.jpg", solid_image(700, 500, 6))
    write_text(
        case / "input.txt",
        'Ampersands & <angles> and "quotes"\n'
        "\n"
        'A paragraph with <b>markup</b>, an & ampersand, and "double quotes".\n'
        "\n"
        "[image:1]\n",
    )


def case_whitespace() -> None:
    """The normalisation rules: NBSP, zero-width space, tabs, and blank-line runs."""
    case = reset_case("whitespace")
    write_image(case / "images" / "photo1.jpg", solid_image(700, 500, 7))
    write_text(
        case / "input.txt",
        "Whitespace  Normalisation\n"
        "\n"
        "\n"
        "\n"
        "A paragraph​ with\tspacing   oddities.\n"
        "A continuation line joined into the same paragraph.\n"
        "\n"
        "[image:1]\n",
    )


def case_malformed_markers() -> None:
    """Marker-shaped text the pattern does not match, which stays as literal text."""
    case = reset_case("malformed-markers")
    write_image(case / "images" / "photo1.jpg", solid_image(700, 500, 3))
    write_text(
        case / "input.txt",
        "Malformed Markers\n"
        "\n"
        "Not a marker: [image:] and [picture:1] and [image 1] and [IMAGE:1].\n"
        "\n"
        "Spaced but valid: [images: 1 , 1 ]\n"
        "\n"
        "[image:1]\n",
    )


# --------------------------------------------------------------------------------------
# Word cases
# --------------------------------------------------------------------------------------

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

# A fixed instant for every generated package, so a .docx depends only on its content.
FIXED_ZIP_TIME = (1980, 1, 1, 0, 0, 0)
FIXED_TIMESTAMP = "2026-01-01T00:00:00Z"


def normalize_docx(path: Path) -> None:
    """Strips the wall-clock time out of a saved package.

    ``python-docx`` stamps the zip entries with the current time and writes
    ``dcterms:created``/``dcterms:modified`` into ``docProps/core.xml``. Both make the file
    differ between runs, which would put a spurious diff in front of every reviewer and break
    the determinism constitution IV asks of the corpus.
    """
    import re as _re

    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        entries = {name: archive.read(name) for name in names}

    core = entries.get("docProps/core.xml")
    if core is not None:
        text = core.decode("utf-8")
        text = _re.sub(
            r"(<dcterms:(?:created|modified)[^>]*>)[^<]*(</dcterms:(?:created|modified)>)",
            rf"\g<1>{FIXED_TIMESTAMP}\g<2>",
            text,
        )
        entries["docProps/core.xml"] = text.encode("utf-8")

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name in names:
            info = zipfile.ZipInfo(name, date_time=FIXED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            archive.writestr(info, entries[name])


def build_docx(path: Path, title: str, paragraphs: list[str], images: list[Path],
               title_style: str | None = "Heading 1") -> None:
    """Writes a .docx whose images sit inline between the paragraphs."""
    from docx import Document
    from docx.shared import Inches

    document = Document()
    heading = document.add_paragraph(title)
    if title_style:
        heading.style = document.styles[title_style]

    image_iter = iter(images)
    for index, text in enumerate(paragraphs):
        document.add_paragraph(text)
        # An image after every second paragraph, which is what gives after_paragraph
        # a value worth asserting.
        if index % 2 == 1:
            picture = next(image_iter, None)
            if picture is not None:
                document.add_picture(str(picture), width=Inches(4.0))
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    normalize_docx(path)


def case_word_with_photos() -> None:
    """A Word document with inline images, the FR-002/FR-003 path."""
    case = reset_case("word-with-photos")
    images = []
    for idx in range(1, 5):
        target = case / "images" / f"photo{idx}.jpg"
        write_image(target, solid_image(900, 600, idx))
        images.append(target)
    build_docx(
        case / "input.docx",
        "Word With Photos",
        [
            "The lead paragraph of a Word document.",
            "A second paragraph, after which the first image sits.",
            "A third paragraph continues the story.",
            "A fourth paragraph, followed by the second image.",
            "A fifth paragraph.",
            "A sixth paragraph, followed by the third image.",
            "A closing paragraph.",
        ],
        images[:3],
    )


def case_word_anchored() -> None:
    """A Word document whose drawings are floating anchors rather than inline runs.

    ``python-docx`` only writes inline drawings, so the anchors are produced by rewriting
    ``wp:inline`` into ``wp:anchor`` in the saved package. The result is what an editor gets
    from Word's "Square" or "Tight" text wrapping, which is the shape ``core::import::docx``
    has to survive.
    """
    case = reset_case("word-anchored")
    images = []
    for idx in range(1, 4):
        target = case / "images" / f"photo{idx}.jpg"
        write_image(target, solid_image(800, 600, idx + 2))
        images.append(target)
    docx_path = case / "input.docx"
    build_docx(
        docx_path,
        "Word Anchored",
        [
            "The lead paragraph of an anchored document.",
            "A second paragraph, with a floating image attached to it.",
            "A third paragraph.",
            "A fourth paragraph, with the second floating image.",
            "A closing paragraph.",
        ],
        images[:2],
    )
    _inline_to_anchor(docx_path)


def _inline_to_anchor(docx_path: Path) -> None:
    """Rewrites every ``wp:inline`` drawing in a package into a floating ``wp:anchor``."""
    from lxml import etree

    wp = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"

    with zipfile.ZipFile(docx_path) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}

    root = etree.fromstring(entries["word/document.xml"])
    for inline in root.iter(f"{{{wp}}}inline"):
        anchor = etree.Element(f"{{{wp}}}anchor", nsmap=inline.nsmap)
        anchor.set("distT", "0")
        anchor.set("distB", "0")
        anchor.set("distL", "114300")
        anchor.set("distR", "114300")
        anchor.set("simplePos", "0")
        anchor.set("relativeHeight", "251658240")
        anchor.set("behindDoc", "0")
        anchor.set("locked", "0")
        anchor.set("layoutInCell", "1")
        anchor.set("allowOverlap", "1")

        simple_pos = etree.SubElement(anchor, f"{{{wp}}}simplePos")
        simple_pos.set("x", "0")
        simple_pos.set("y", "0")
        for tag, relative, offset in (("positionH", "column", "0"), ("positionV", "paragraph", "0")):
            node = etree.SubElement(anchor, f"{{{wp}}}{tag}")
            node.set("relativeFrom", relative)
            child = etree.SubElement(node, f"{{{wp}}}posOffset")
            child.text = offset
        # Carry the inline children across unchanged: extent, docPr, and the graphic itself.
        for child in list(inline):
            anchor.append(copy.deepcopy(child))
        etree.SubElement(anchor, f"{{{wp}}}wrapSquare").set("wrapText", "bothSides")

        parent = inline.getparent()
        parent.replace(inline, anchor)

    entries["word/document.xml"] = etree.tostring(
        root, xml_declaration=True, encoding="UTF-8", standalone=True
    )

    with zipfile.ZipFile(docx_path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, payload in entries.items():
            info = zipfile.ZipInfo(name, date_time=FIXED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            archive.writestr(info, payload)


def case_word_plain_title() -> None:
    """A Word document with no Heading style: the first paragraph becomes the title."""
    case = reset_case("word-plain-title")
    write_image(case / "images" / "photo1.jpg", solid_image(800, 600, 1))
    build_docx(
        case / "input.docx",
        "A Plain First Paragraph As Title",
        ["The body starts here.", "And continues here.", "And ends here."],
        [],
        title_style=None,
    )


def case_word_unsupported() -> None:
    """A Word document carrying a table, which FR-034 says to skip with a warning."""
    case = reset_case("word-unsupported")
    write_image(case / "images" / "photo1.jpg", solid_image(800, 600, 2))
    from docx import Document

    document = Document()
    document.add_paragraph("Word Unsupported").style = document.styles["Heading 1"]
    document.add_paragraph("A paragraph before the table.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Cell A"
    table.cell(0, 1).text = "Cell B"
    table.cell(1, 0).text = "Cell C"
    table.cell(1, 1).text = "Cell D"
    document.add_paragraph("A paragraph after the table.")
    (case / "input.docx").parent.mkdir(parents=True, exist_ok=True)
    document.save(str(case / "input.docx"))
    normalize_docx(case / "input.docx")


CASES = [
    case_markers,
    case_single_markers,
    case_row_of_one,
    case_cyrillic_title,
    case_cyrillic_variants,
    case_natural_sort,
    case_missing_photo,
    case_unused_photo,
    case_oversized,
    case_formats,
    case_mixed_orientation,
    case_portraits,
    case_markdown_title,
    case_escaping,
    case_whitespace,
    case_malformed_markers,
    case_word_with_photos,
    case_word_anchored,
    case_word_plain_title,
    case_word_unsupported,
]


def main() -> int:
    INPUTS.mkdir(parents=True, exist_ok=True)
    for case in CASES:
        case()
        print(f"generated {case.__name__.removeprefix('case_')}")
    kept = sorted(name for name in PRESERVED if (INPUTS / name).exists())
    print(f"preserved real-document cases: {', '.join(kept) or 'none'}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
